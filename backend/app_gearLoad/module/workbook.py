"""齿轮载荷导出的 xlsx：每级一张表，写法同网页版在浏览器里用 xlsx 库生成的

表头和数据行就是 `CalGearLoads.getExportData` 给的，原样写进表：第一列是工况序号（字符串，从 0 起），
算不出来的格子是空串；序号列宽 8，其余列宽 15。
"""

from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

#: 序号列和数值列的列宽（字符数），同网页版
INDEX_WIDTH = 8
VALUE_WIDTH = 15


def write_workbook(tables: dict[str, dict[str, Any]], path: str) -> None:
    """把各级的导出结果写成一个 xlsx，已有的同名文件直接覆盖。

    Args:
        tables: 级名（stageI……）→ `{"headers": 表头, "data": 数据行}`，一级一张表，表名就是级名。
        path: xlsx 的完整路径。
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, table in tables.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(table["headers"])
        for row in table["data"]:
            sheet.append(row)
        for index in range(len(table["headers"])):
            width = INDEX_WIDTH if index == 0 else VALUE_WIDTH
            sheet.column_dimensions[get_column_letter(index + 1)].width = width
    workbook.save(path)
