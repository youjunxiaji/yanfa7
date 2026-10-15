"""app_gearLoad/views.py 的 calculate、export_excel —— 照搬网页版的计算和导出 Excel

期望值是拿网页版原代码（tmb_web_v1 的 shanghai_yanfa7/utils.py）算同一组输入得到的，保证搬过来算得一样。
输入表照网页版页面发的写：16 行，一行一个英文参数名，4 级的值都是字符串。
"""

from openpyxl import load_workbook

from app_gearLoad import views

PARAMS = [
    "Gear stage arrangement",
    "Nominal pressure angle αn (°)",
    "Helix angle β (°)",
    "Helix angle LH of RH - sun wheel or pinion",
    "Module in the normal section mn (mm)",
    "Number of teeth ring wheel zR / wheel z1",
    "Number of teeth planet zpl / pinion z2",
    "Number of teeth sun zsun",
    "Centre distance a (mm)",
    "Number of planets",
    "Power P (kW)",
    "Torque T (kNm)",
    "Input speed n (RPM)",
    "Root diameter df (mm)",
    "Tip diameter da (mm)",
    "Load factor KY",
]
STAGES = ["stageI", "stageII", "stageIII", "stageIIII"]

# 没填的一级：网页版页面上布置默认 Planetary、旋向默认 L，其余空着
EMPTY = ["Planetary", "", "", "L"] + [""] * 12
# 行星级：αn 20°、β 7°、mn 12、zR 99 / zpl 36 / zsun 27、a 381、3 个行星轮、P 5000、T 1000、n 12、df / da 0、KY 1.1
PLANETARY = ["Planetary", "20", "7", "L", "12", "99", "36", "27", "381", "3", "5000", "1000", "12", "0", "0", "1.1"]
# 平行级：用不到的太阳轮齿数、行星轮数填 0，KY 填 1；扭矩、转速随便填，算的时候取上一级的输出
HELICAL = ["Helical", "20", "10", "R", "10", "98", "23", "0", "610", "0", "5000", "1", "1", "260", "290", "1"]

# 网页版原代码算出的结果
STAGE_I = {
    "input_data": [5000.0, 214.286, 56.0],
    "gear_loads": [962.38, 481.19, 177.084, 59.108, 481.19, 177.084, 59.108, 214.286, 161.203],
    "speed_of_components": [56.0, -33.0, 0],
    "equivalent_diameter": [217.714, 217.714, ""],
}
STAGE_II = {
    "input_data": [5000.0, 50.292, 238.609],
    "gear_loads": ["", 433.734, 150.635, 75.94, "", "", "", 50.292, ""],
    "speed_of_components": [56.0, 238.609, ""],
    "equivalent_diameter": [275.0, "", ""],
}

EXPORT_HEADERS = [
    "切向力 2Ftan(KN)",
    "行星轮-太阳轮-切向力(KN)",
    "行星轮-太阳轮-径向力(KN)",
    "行星轮-太阳轮-轴向力(KN)",
    "行星轮-齿圈-切向力(KN)",
    "行星轮-齿圈-径向力(KN)",
    "行星轮-齿圈-轴向力(KN)",
    "扭矩T(KNm)",
    "轴向力太阳轮(KN)",
    "太阳轮转速(RPM)",
    "行星轮转速(RPM)",
    "齿圈转速(RPM)",
]

# 两个工况：T 1000 / n 12，T 800 / n 10.5。键的顺序照网页版新录入时的写法，扭矩在前
CONDITIONS = [{"torque": "1000", "rpm": "12"}, {"torque": "800", "rpm": "10.5"}]


def input_data(*stages: list[str]) -> list[dict]:
    """网页版页面发的输入表，没给的级空着。"""
    columns = list(stages) + [EMPTY] * (4 - len(stages))
    return [
        {
            "parameter": name,
            **{stage: column[index] for stage, column in zip(STAGES, columns)},
            "editable": index not in (0, 3),
        }
        for index, name in enumerate(PARAMS)
    ]


def calculate(api_client, *stages: list[str]):
    return api_client.post("/api/gear-load/calculate", json={"inputData": input_data(*stages)})


def export(api_client, path, **overrides):
    body = {"inputData": input_data(PLANETARY, HELICAL), "selectedData": CONDITIONS, "path": str(path)}
    body.update(overrides)
    return api_client.post("/api/gear-load/export", json=body)


def sheet_rows(path, sheet: str) -> list[list]:
    """读回一张表，空格子读出来是 None。"""
    return [list(row) for row in load_workbook(path)[sheet].iter_rows(values_only=True)]


class TestCalculate:
    def test_一级行星_和网页版算得一样(self, api_client):
        response = calculate(api_client, PLANETARY)

        assert response.status_code == 200
        assert response.json() == {"stageI": STAGE_I}

    def test_二级的扭矩和转速取一级的输出(self, api_client):
        response = calculate(api_client, PLANETARY, HELICAL)

        assert response.json() == {"stageI": STAGE_I, "stageII": STAGE_II}

    def test_只算整列填满的级(self, api_client):
        no_distance = HELICAL.copy()
        no_distance[PARAMS.index("Centre distance a (mm)")] = ""

        response = calculate(api_client, PLANETARY, no_distance)

        assert response.json() == {"stageI": STAGE_I}

    def test_一级没填满什么都不算(self, api_client):
        no_power = PLANETARY.copy()
        no_power[PARAMS.index("Power P (kW)")] = ""

        response = calculate(api_client, no_power)

        assert response.json() == {}


class TestExport:
    def test_每级一张表_每个工况一行(self, api_client, tmp_path):
        path = tmp_path / "齿轮载荷.xlsx"

        response = export(api_client, path)

        assert response.status_code == 200
        assert response.json() == {"path": str(path)}
        assert load_workbook(path).sheetnames == ["stageI", "stageII"]
        assert sheet_rows(path, "stageI") == [
            [None, *EXPORT_HEADERS],
            ["0", 962.38, 481.19, 177.084, 59.108, 481.19, 177.084, 59.108, 214.286, 161.203, 56, -33, 0],
            ["1", 769.904, 384.952, 141.667, 47.286, 384.952, 141.667, 47.286, 171.429, 128.962, 49, -28.875, 0],
        ]
        assert sheet_rows(path, "stageII") == [
            [None, *EXPORT_HEADERS],
            ["0", None, 433.734, 150.635, 75.94, None, None, None, 50.292, None, 56, 238.609, None],
            ["1", None, 346.987, 120.508, 60.752, None, None, None, 40.233, None, 49, 208.783, None],
        ]

    def test_第一个工况和页面上算的一样(self, api_client, tmp_path):
        path = tmp_path / "out.xlsx"

        export(api_client, path)

        # 网页版新录入的工况在这里会把扭矩、转速对调
        first = sheet_rows(path, "stageI")[1]
        assert first[1:10] == STAGE_I["gear_loads"]
        assert first[10:] == STAGE_I["speed_of_components"]

    def test_列宽同网页版(self, api_client, tmp_path):
        path = tmp_path / "out.xlsx"

        export(api_client, path)

        sheet = load_workbook(path)["stageI"]
        assert sheet.column_dimensions["A"].width == 8
        assert {sheet.column_dimensions[column].width for column in "BCDEFGHIJKLM"} == {15}

    def test_没有多工况_400(self, api_client, tmp_path):
        response = export(api_client, tmp_path / "out.xlsx", selectedData=[])

        assert response.status_code == 400
        assert "多工况" in response.json()["detail"]

    def test_一级没填满_400(self, api_client, tmp_path):
        no_power = PLANETARY.copy()
        no_power[PARAMS.index("Power P (kW)")] = ""

        response = export(api_client, tmp_path / "out.xlsx", inputData=input_data(no_power))

        assert response.status_code == 400
        assert "Stage I" in response.json()["detail"]

    def test_文件被占用_400(self, api_client, tmp_path, monkeypatch):
        def locked(tables, path):
            raise PermissionError(13, "Permission denied", path)

        monkeypatch.setattr(views, "write_workbook", locked)

        response = export(api_client, tmp_path / "out.xlsx")

        assert response.status_code == 400
        assert response.json()["detail"] == "out.xlsx 正被其他程序占用，请关闭后重试"

    def test_文件夹不存在_400(self, api_client, tmp_path):
        response = export(api_client, tmp_path / "没有这个文件夹" / "out.xlsx")

        assert response.status_code == 400
        assert response.json()["detail"].startswith("保存 Excel 失败")
