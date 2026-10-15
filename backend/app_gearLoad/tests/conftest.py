"""齿轮载荷测试的 fixture；假 TMB 服务器 `tmb_server` 几个 app 共用，在 backend/conftest.py

`calculate`、`export_tables` 是用例调生产代码的唯一入口。优化算法、改接口的写法时只改这两个 fixture，
test_calculation.py、test_export.py 里的用例和断言都不动。
"""

import itertools
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app_gearLoad.tests.factories import Gearbox, input_rows, selected_rows


@pytest.fixture
def api_client():
    """FastAPI 测试客户端：进程内调应用，不开端口。"""
    from fastapi.testclient import TestClient

    import main

    with TestClient(main.app) as client:
        yield client


@pytest.fixture
def calculate(api_client):
    """用生产代码算一个工况，结果同网页版：级名 → input_data、gear_loads、speed_of_components、equivalent_diameter。"""

    def run(gearbox: Gearbox, torque: float, speed: float) -> dict:
        response = api_client.post(
            "/api/gear-load/calculate", json={"inputData": input_rows(gearbox, torque, speed)}
        )
        assert response.status_code == 200, response.text
        return response.json()

    return run


@pytest.fixture
def export_tables(api_client, tmp_path):
    """用生产代码导出多工况，把写出的 xlsx 读回来：表名 → {"headers": 表头, "data": 数据行}。

    页面上 Stage I 的扭矩、转速是第一组工况，同网页版点「确定」以后的样子。
    """
    counter = itertools.count()

    def run(gearbox: Gearbox, conditions) -> dict:
        path = tmp_path / f"export-{next(counter)}.xlsx"
        body = {
            "inputData": input_rows(gearbox, *conditions[0]),
            "selectedData": selected_rows(conditions),
            "path": str(path),
        }
        response = api_client.post("/api/gear-load/export", json=body)
        assert response.status_code == 200, response.text
        return read_tables(path)

    return run


def read_tables(path: Path) -> dict:
    workbook = load_workbook(path)
    tables = {}
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        tables[sheet.title] = {"headers": rows[0], "data": rows[1:]}
    return tables
