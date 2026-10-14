"""真实算例全流程对拍 —— 基线是并入 yanfa7 之前 simple_load v1.3.1 的输出

`测试案例/基线_v1.3.1/` 不入库，本地没有就整体跳过。基线由同目录的 `生成基线.py`
在 simple_load 仓库里跑出来：GW-V16 的 300 个文件，扭矩开关关、开各一次。
这里用 `params.json` 里的同一套参数跑现在的代码，接口返回值和两个 Excel 逐格比对：
数值按相对误差 `REL_TOL` 比，文字、工作表顺序和合并单元格要求完全一致。

迁移期间这条用例是判卷标准：它红了说明计算结果变了，先查清楚，不改基线。
"""

import json
import math
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app_simpleLoad.module.config import AxisMapping, ConversionConfig, PathConfig
from app_simpleLoad.module.cal_simpleLoad import CalSimpleLoad

BASELINE = "基线_v1.3.1"
RUNS = ["扭矩关", "扭矩开"]

#: polars 多线程分组求和的顺序不固定，同一份代码连跑两次，缩减结果的末位就不一样
#: （实测和基线的相对差不超过 2e-15）。1e-9 远大于这种误差，又比 float32 的精度（约 1e-7）严，
#: 哪一步把计算精度改了能抓出来。
REL_TOL = 1e-9
#: 0 只和 0 相等：基线里为 0 的格子都是确定算出来的 0
ABS_TOL = 0.0


@pytest.fixture(scope="module")
def baseline_dir(real_case_dir: Path) -> Path:
    folder = real_case_dir / BASELINE
    if not (folder / "params.json").is_file():
        pytest.skip(f"本地没有 测试案例/{BASELINE}/，跳过基线对拍")
    return folder


def is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def assert_close(actual, expected, where: str) -> None:
    """数值按容差比，其余（文字、None、布尔）要求相等。"""
    if is_number(actual) and is_number(expected):
        assert math.isclose(actual, expected, rel_tol=REL_TOL, abs_tol=ABS_TOL), \
            f"{where}：{actual!r} ≠ 基线 {expected!r}"
    else:
        assert actual == expected, f"{where}：{actual!r} ≠ 基线 {expected!r}"


def assert_close_json(actual, expected, where: str) -> None:
    """逐层比对 JSON 解析出来的结构：键集合、列表长度一致，叶子按 `assert_close` 比。"""
    if isinstance(expected, dict):
        assert isinstance(actual, dict) and set(actual) == set(expected), f"{where}：键不同"
        for key, value in expected.items():
            assert_close_json(actual[key], value, f"{where}.{key}")
    elif isinstance(expected, list):
        assert isinstance(actual, list) and len(actual) == len(expected), f"{where}：长度不同"
        for index, value in enumerate(expected):
            assert_close_json(actual[index], value, f"{where}[{index}]")
    else:
        assert_close(actual, expected, where)


def assert_same_workbook(actual_path: Path, expected_path: Path) -> None:
    """两个 Excel 逐个工作表、逐格比对，合并单元格的范围也要一致。"""
    actual_book = load_workbook(actual_path)
    expected_book = load_workbook(expected_path)
    assert actual_book.sheetnames == expected_book.sheetnames, f"{actual_path.name}：工作表不同"

    for name in expected_book.sheetnames:
        actual_sheet = actual_book[name]
        expected_sheet = expected_book[name]
        where = f"{actual_path.name} / {name}"

        assert sorted(map(str, actual_sheet.merged_cells.ranges)) == \
            sorted(map(str, expected_sheet.merged_cells.ranges)), f"{where}：合并单元格不同"
        assert (actual_sheet.max_row, actual_sheet.max_column) == \
            (expected_sheet.max_row, expected_sheet.max_column), f"{where}：行列数不同"

        rows = zip(actual_sheet.iter_rows(values_only=True), expected_sheet.iter_rows(values_only=True))
        for row_no, (actual_row, expected_row) in enumerate(rows, start=1):
            for col_no, (actual, expected) in enumerate(zip(actual_row, expected_row), start=1):
                assert_close(actual, expected, f"{where} 第 {row_no} 行第 {col_no} 列")


@pytest.mark.parametrize("run", RUNS)
async def test_与_v1_3_1_的输出一致(baseline_dir: Path, real_case_dir: Path, tmp_path: Path, run: str):
    params = json.loads((baseline_dir / "params.json").read_text(encoding="utf-8"))
    spec = params["runs"][run]

    # 输出文件名取保存目录的名字，和基线同名才好一一对上
    out_dir = tmp_path / run
    out_dir.mkdir()

    instance = CalSimpleLoad()
    instance.setInit(
        paths=PathConfig(
            result_folder_save_path=str(out_dir),
            load_file_folder_path=str(real_case_dir / params["load_file_folder"]),
            freq_table_path=str(real_case_dir / params["freq_table"]),
            keep_torque_component=spec["keep_torque_component"],
        ),
        header=params["header"],
        config=ConversionConfig(**params["conversion_factors"]),
    )
    romax_origin = [AxisMapping(**item) for item in params["romax_origin"]]

    await instance.simple_Pre_processing()
    min_max = await instance.simple_load1(romax_origin)
    echarts_data = await instance.savePic()
    count = await instance.simple_load2(spec["tableData"], romax_origin)

    expected = json.loads((baseline_dir / run / "接口返回.json").read_text(encoding="utf-8"))
    assert count == expected["count"]
    assert_close_json(json.loads(min_max), json.loads(expected["min_max"]), "min_max")
    assert set(echarts_data) == set(expected["echarts_data"])
    for column, series in expected["echarts_data"].items():
        assert_close_json(json.loads(echarts_data[column]), json.loads(series), f"echarts_data.{column}")

    for kind in ("GL", "Romax"):
        name = f"Load_Reduction_{kind}-{run}.xlsx"
        assert_same_workbook(out_dir / name, baseline_dir / run / name)
